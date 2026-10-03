from __future__ import annotations
from copy import deepcopy
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

def _parse_alignment(a):
    if not a: return None
    if a.horizontal in {'center','centerContinuous','distributed','fill','justify'}: return 'center'
    return a.horizontal if a.horizontal in {'left','right'} else None

def _parse_number_format(fmt):
    if not fmt or fmt.lower() in {'general','@'}: return None
    lower=fmt.lower(); dec=fmt.split('.')[-1].count('0') if '.' in fmt else 0
    if '%' in fmt: return {'type':'percent','decimals':dec,'thousands':',' in fmt}
    if '¥' in fmt or '[$¥' in fmt: return {'type':'currency','decimals':dec,'thousands':',' in fmt}
    if any(x in lower for x in ['yy','dd']): return {'type':'date','decimals':0,'thousands':False}
    if any(x in lower for x in ['hh','ss']): return {'type':'time','decimals':0,'thousands':False}
    if any(x in fmt for x in '0#'): return {'type':'number','decimals':dec,'thousands':',' in fmt}
    return None

def load_xlsx(file_path: str) -> dict[str, dict]:
    wb=load_workbook(file_path,data_only=False); result={}
    for ws in wb.worksheets:
        max_row=max(ws.max_row or 1,1); max_col=max(ws.max_column or 1,1); rows=[]; styles={}
        for r in range(1,max_row+1):
            vals=[]
            for c in range(1,max_col+1):
                cell=ws.cell(r,c); vals.append('' if cell.value is None else str(cell.value)); st={}
                if cell.font:
                    if cell.font.bold: st['bold']=True
                    if cell.font.italic: st['italic']=True
                    if cell.font.underline and cell.font.underline!='none': st['underline']=True
                    if cell.font.sz: st['font_size']=int(cell.font.sz)
                    if cell.font.color and getattr(cell.font.color,'type',None)=='rgb' and cell.font.color.rgb: st['fg']='#'+cell.font.color.rgb[-6:]
                if cell.fill and cell.fill.fill_type=='solid' and getattr(cell.fill.fgColor,'type',None)=='rgb' and cell.fill.fgColor.rgb: st['bg']='#'+cell.fill.fgColor.rgb[-6:]
                ha=_parse_alignment(cell.alignment)
                if ha: st['halign']=ha
                borders={}
                if cell.border:
                    for side in ['left','top','right','bottom']:
                        if getattr(cell.border,side) and getattr(cell.border,side).style: borders[side]=True
                if borders: st['borders']=borders
                nf=_parse_number_format(cell.number_format)
                if nf: st['number_format']=nf
                if st: styles[f'{r-1},{c-1}']=deepcopy(st)
            rows.append(vals)
        result[ws.title]={'rows':max_row,'cols':max_col,'data':rows,'styles':styles,
            'view':{'column_widths':{str(c-1):ws.column_dimensions[get_column_letter(c)].width for c in range(1,max_col+1) if ws.column_dimensions[get_column_letter(c)].width},
                    'row_heights':{str(r-1):ws.row_dimensions[r].height for r in range(1,max_row+1) if ws.row_dimensions[r].height},
                    'hidden_columns':[c-1 for c in range(1,max_col+1) if ws.column_dimensions[get_column_letter(c)].hidden],
                    'hidden_rows':[r-1 for r in range(1,max_row+1) if ws.row_dimensions[r].hidden],
                    'freeze_panes':str(ws.freeze_panes) if ws.freeze_panes else None,
                    'merged_cells':[str(x) for x in ws.merged_cells.ranges]}}
    return result

def _apply_style(cell, st):
    kw={}
    for k in ['bold','italic']:
        if st.get(k): kw[k]=True
    if st.get('underline'): kw['underline']='single'
    if st.get('font_size'): kw['sz']=st['font_size']
    if st.get('fg'): kw['color']=st['fg'].replace('#','')
    if kw: cell.font=Font(**kw)
    if st.get('bg'): cell.fill=PatternFill(fill_type='solid',fgColor=st['bg'].replace('#',''))
    if st.get('halign'): cell.alignment=Alignment(horizontal=st['halign'])
    b=st.get('borders') or {}; thin=Side(style='thin',color='000000')
    if b: cell.border=Border(**{x:(thin if b.get(x) else Side(style=None)) for x in ['left','right','top','bottom']})
    nf=st.get('number_format')
    if isinstance(nf,dict):
        t=nf.get('type','general'); d=int(nf.get('decimals',2)); th=bool(nf.get('thousands')); base=('#,##0' if th else '0')+('.'+'0'*d if d else '')
        if t=='number': cell.number_format=base
        elif t=='percent': cell.number_format=base+'%'
        elif t=='currency': cell.number_format='[$¥-411]'+base
        elif t=='date': cell.number_format='yyyy-mm-dd'
        elif t=='time': cell.number_format='hh:mm:ss'

def save_xlsx(file_path: str, sheets: dict[str, dict|list[list[str]]]) -> None:
    wb=Workbook(); first=True
    for name,p in sheets.items():
        ws=wb.active if first else wb.create_sheet(); ws.title=name; first=False
        matrix=p if isinstance(p,list) else p.get('data',[]); styles={} if isinstance(p,list) else p.get('styles',{}); view={} if isinstance(p,list) else p.get('view',{})
        for r,row in enumerate(matrix,1):
            for c,value in enumerate(row,1):
                cell=ws.cell(r,c); cell.value=value
                if styles.get(f'{r-1},{c-1}'): _apply_style(cell,styles[f'{r-1},{c-1}'])
        for c,w in view.get('column_widths',{}).items(): ws.column_dimensions[get_column_letter(int(c)+1)].width=float(w)
        for r,h in view.get('row_heights',{}).items(): ws.row_dimensions[int(r)+1].height=float(h)
        for c in view.get('hidden_columns',[]): ws.column_dimensions[get_column_letter(int(c)+1)].hidden=True
        for r in view.get('hidden_rows',[]): ws.row_dimensions[int(r)+1].hidden=True
        if view.get('freeze_panes'): ws.freeze_panes=view['freeze_panes']
        elif view.get('freeze_top_row') and view.get('freeze_first_col'): ws.freeze_panes='B2'
        elif view.get('freeze_top_row'): ws.freeze_panes='A2'
        elif view.get('freeze_first_col'): ws.freeze_panes='B1'
        for rng in view.get('merged_cells',[]):
            try: ws.merge_cells(rng)
            except ValueError: pass
    wb.save(file_path)
